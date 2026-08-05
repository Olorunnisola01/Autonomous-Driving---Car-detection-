#!/usr/bin/env python3
"""Regenerate MASTER_DOCUMENT.docx from the editable Markdown and JSON sources.

Requires python-docx. Example:
    ../.venv/bin/python build_master.py
"""
from __future__ import annotations

import json
from pathlib import Path

try:
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_BREAK
    from docx.shared import Inches, Pt, RGBColor
except ImportError as exc:
    raise SystemExit("python-docx is required. Install it in an environment, then rerun build_master.py.") from exc

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "MASTER_DOCUMENT.docx"


def prepare_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(4)
    for style_name, size, colour in (("Title", 25, RGBColor(16, 44, 58)), ("Heading 1", 17, RGBColor(16, 44, 58)), ("Heading 2", 13, RGBColor(25, 99, 107))):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = colour
    code = doc.styles.add_style("Prompt Code", WD_STYLE_TYPE.PARAGRAPH) if "Prompt Code" not in doc.styles else doc.styles["Prompt Code"]
    code.font.name = "Courier New"
    code.font.size = Pt(7.5)
    code.paragraph_format.space_after = Pt(3)
    return doc


def add_markdown(doc: Document, path: Path) -> None:
    """Render readable Markdown prose without pretending to be a full Markdown engine."""
    lines = path.read_text(encoding="utf-8").splitlines()
    paragraph = []
    def flush() -> None:
        if paragraph:
            doc.add_paragraph(" ".join(x.strip() for x in paragraph))
            paragraph.clear()
    for raw in lines:
        text = raw.strip()
        if not text:
            flush(); continue
        if text.startswith("# "):
            flush(); doc.add_heading(text[2:], level=1)
        elif text.startswith("## "):
            flush(); doc.add_heading(text[3:], level=2)
        elif text.startswith("### "):
            flush(); doc.add_heading(text[4:], level=3)
        elif text.startswith("#### "):
            flush(); doc.add_heading(text[5:], level=4)
        elif text.startswith("> "):
            flush(); p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(.25); run = p.add_run(text[2:]); run.italic = True
        elif text.startswith("- "):
            flush(); doc.add_paragraph(text[2:], style="List Bullet")
        elif text.startswith("|"):
            # Preserve the planning tables as compact monospace rows, omitting separator rows.
            if set(text.replace("|", "").strip()) <= {"-", ":"}:
                continue
            flush(); doc.add_paragraph(text, style="Prompt Code")
        else:
            paragraph.append(text)
    flush()


def main() -> None:
    prompts = json.loads((ROOT / "prompts.json").read_text(encoding="utf-8"))
    doc = prepare_document()
    doc.add_heading("THE RIVER REMEMBERS", 0)
    doc.add_paragraph("Master production document · screenplay, continuity bible, location registry and 200 structured image prompts.")
    doc.add_paragraph("Regenerate this file with build_master.py after editing SCREENPLAY.md, LOCATION_REGISTRY.md or prompts.json.")
    doc.add_page_break()
    add_markdown(doc, ROOT / "SCREENPLAY.md")
    doc.add_page_break()
    doc.add_heading("Location Registry", 0)
    add_markdown(doc, ROOT / "LOCATION_REGISTRY.md")
    doc.add_page_break()
    doc.add_heading("Structured Shot Prompts", 0)
    doc.add_paragraph(f"{len(prompts['shots'])} planned shot objects. Wardrobe and environment continuity strings are preserved verbatim below.")
    current_scene = None
    for shot in prompts["shots"]:
        if shot["scene"] != current_scene:
            current_scene = shot["scene"]
            doc.add_heading(current_scene, level=1)
        p = doc.add_paragraph(style="Prompt Code")
        p.add_run(json.dumps(shot, ensure_ascii=False, indent=2))
        p.paragraph_format.keep_together = True
    doc.save(OUT)
    print(f"Wrote {OUT.name} with screenplay, location registry and {len(prompts['shots'])} JSON prompt blocks.")

if __name__ == "__main__":
    main()
